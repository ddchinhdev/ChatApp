from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = r"E:\Chatapp\ChatApp\Phan_cong_backend_3_thanh_vien_3_tuan.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=130, start=150, bottom=130, end=150):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_bullet(cell, text):
    p = cell.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(9.2)


doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width
section.top_margin = Cm(1.6)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(1.5)
section.right_margin = Cm(1.5)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10)
styles["Title"].font.name = "Aptos Display"
styles["Title"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
styles["Title"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].font.size = Pt(22)
styles["Title"].font.bold = True

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Phân công Backend cho 3 thành viên trong 3 tuần")
title.paragraph_format.space_after = Pt(8)

intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
intro.paragraph_format.space_after = Pt(12)
run = intro.add_run(
    "Phân chia theo module để hạn chế xung đột code. Thành viên 1 đảm nhận khối lượng lớn hơn; "
    "Thành viên 2 và Thành viên 3 có khối lượng tương đương."
)
run.font.size = Pt(10.5)
run.font.color.rgb = RGBColor(64, 64, 64)

table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.style = "Table Grid"

headers = [
    "Tiến độ",
    "Thành viên 1\nAuth User và Group",
    "Thành viên 2\nConversation và Message",
    "Thành viên 3\nRealtime và Admin",
]
widths = [2.5, 8.0, 8.0, 8.0]
for idx, cell in enumerate(table.rows[0].cells):
    set_cell_width(cell, widths[idx])
    set_cell_shading(cell, "17365D")
    set_cell_margins(cell, 150, 150, 150, 150)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(headers[idx])
    r.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    r.font.size = Pt(10)
set_repeat_table_header(table.rows[0])

weeks = [
    (
        "Tuần 1\nXây nền",
        [
            "Khởi tạo Spring Boot; cấu hình MySQL, JPA và Flyway.",
            "Thiết kế users và roles; DTO, repository và validation.",
            "Đăng ký, đăng nhập; JWT; Spring Security; phân quyền USER và ADMIN.",
        ],
        [
            "Thiết kế conversations, conversation_members và messages.",
            "Tạo hoặc lấy conversation 1-1; chống tạo trùng.",
            "Danh sách conversation; kiểm tra quyền thành viên.",
        ],
        [
            "Thiết kế WebSocket và STOMP; thống nhất event và payload.",
            "Cấu hình endpoint /ws và chuẩn bị JWT WebSocket.",
            "Chuẩn bị Swagger, CORS và Docker database.",
        ],
    ),
    (
        "Tuần 2\nChức năng chính",
        [
            "Cập nhật hồ sơ và tìm kiếm người dùng.",
            "Tạo nhóm; cập nhật tên nhóm; thêm và xóa thành viên.",
            "Phân quyền OWNER, ADMIN, MEMBER; kiểm tra quyền quản trị nhóm.",
        ],
        [
            "Gửi và lưu message; cung cấp REST fallback.",
            "Cập nhật last message; phân trang lịch sử bằng cursor.",
            "Dùng clientMessageId để chống message trùng; kiểm tra quyền gửi và đọc.",
        ],
        [
            "JWT WebSocket Authentication; gửi và nhận realtime.",
            "Online và offline; heartbeat; quản lý nhiều session.",
            "Reconnect và cập nhật lastSeenAt.",
        ],
    ),
    (
        "Tuần 3\nHoàn thiện",
        [
            "Chuyển owner; nâng hoặc hạ role; rời hoặc giải tán nhóm.",
            "Phát system message khi thông tin nhóm thay đổi.",
            "Global exception, kiểm thử authorization, tích hợp module và hỗ trợ frontend.",
        ],
        [
            "Unread count; trạng thái SENT, DELIVERED và READ; mark as read.",
            "Typing indicator và last seen.",
            "Delete hoặc recall nếu còn thời gian; integration test cho message.",
        ],
        [
            "Admin API: danh sách user, khóa hoặc mở khóa và thống kê.",
            "Audit log; integration test và kiểm tra quyền riêng tư.",
            "Hoàn thiện Docker, tài liệu và kiểm thử realtime.",
        ],
    ),
]

for row_idx, (week, *assignments) in enumerate(weeks, start=1):
    cells = table.add_row().cells
    bg = "F3F7FB" if row_idx % 2 else "FFFFFF"
    for idx, cell in enumerate(cells):
        set_cell_width(cell, widths[idx])
        set_cell_margins(cell)
        set_cell_shading(cell, bg)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cells[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(week)
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(23, 54, 93)
    for col_idx, items in enumerate(assignments, start=1):
        cells[col_idx].paragraphs[0]._element.getparent().remove(cells[col_idx].paragraphs[0]._element)
        for item in items:
            add_bullet(cells[col_idx], item)

doc.add_paragraph().paragraph_format.space_after = Pt(2)
summary_heading = doc.add_paragraph()
set_keep_with_next(summary_heading)
summary_heading.paragraph_format.space_before = Pt(3)
summary_heading.paragraph_format.space_after = Pt(5)
rh = summary_heading.add_run("Tỷ lệ khối lượng")
rh.bold = True
rh.font.size = Pt(12)
rh.font.color.rgb = RGBColor(0, 0, 0)

summary = doc.add_table(rows=2, cols=3)
summary.alignment = WD_TABLE_ALIGNMENT.CENTER
summary.autofit = False
summary.style = "Table Grid"
summary_data = [
    ["Thành viên 1", "Thành viên 2", "Thành viên 3"],
    ["12 điểm  37,5%", "10 điểm  31,25%", "10 điểm  31,25%"],
]
for i, row in enumerate(summary.rows):
    for j, cell in enumerate(row.cells):
        set_cell_width(cell, 8.8)
        set_cell_margins(cell, 110, 120, 110, 120)
        set_cell_shading(cell, "D9EAF7" if i == 0 else "FFFFFF")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(summary_data[i][j])
        r.font.size = Pt(9.5)
        r.bold = i == 0

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("Kế hoạch phân công Backend ChatApp")
fr.font.size = Pt(8)
fr.font.color.rgb = RGBColor(100, 100, 100)

doc.core_properties.title = "Phân công Backend cho 3 thành viên trong 3 tuần"
doc.core_properties.subject = "Bảng phân công chức năng Backend ChatApp"
doc.save(OUTPUT)
print(OUTPUT)
